"""Export & custom-service management for operators.

- /export/excel  → premium-only download of all customers, services, payments
- /my-services   → list operator's custom services
- /my-services/new, /<id>/edit, /<id>/delete → CRUD with editable price
"""
import io
from datetime import datetime
from functools import wraps

from flask import (
    Blueprint, render_template, redirect, url_for, flash, request, abort, send_file
)
from flask_login import login_required, current_user
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from extensions import db
from models import Customer, Service, Payment, CustomService

export_bp = Blueprint("export", __name__)


def _is_premium(user):
    sub = getattr(user, "subscription", None)
    if not sub:
        return False
    paid = (sub.plan_name or "").lower().startswith("premium")
    return bool(sub.is_active and paid and sub.days_left > 0)


def premium_required(f):
    @wraps(f)
    def w(*a, **kw):
        if not _is_premium(current_user):
            flash("This feature is available for Premium subscribers. Upgrade to download your full history.", "warning")
            return redirect(url_for("premium.plans"))
        return f(*a, **kw)
    return w


# ---------------- Excel Export ----------------

def _style_header(ws, headers):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="4F46E5")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = 22


@export_bp.route("/export/excel")
@login_required
@premium_required
def excel_all():
    wb = Workbook()
    user_id = current_user.id

    # Customers sheet
    ws = wb.active
    ws.title = "Customers"
    _style_header(ws, ["ID", "Full Name", "Mobile", "Aadhaar L4", "Address", "Village", "Notes", "Created"])
    for c in Customer.query.filter_by(user_id=user_id).order_by(Customer.created_at).all():
        ws.append([c.id, c.full_name, c.mobile_number, c.aadhaar_last_4, c.address,
                   c.village, c.notes, c.created_at.strftime("%Y-%m-%d %H:%M") if c.created_at else ""])

    # Services sheet
    ws2 = wb.create_sheet("Services")
    _style_header(ws2, ["ID", "Customer", "Service Type", "Status", "Amount", "Submitted", "Completed", "Notes"])
    services = (Service.query.join(Customer)
                .filter(Customer.user_id == user_id)
                .order_by(Service.submission_date).all())
    for s in services:
        ws2.append([s.id, s.customer.full_name, s.service_type, s.status, float(s.amount or 0),
                    s.submission_date.strftime("%Y-%m-%d %H:%M") if s.submission_date else "",
                    s.completion_date.strftime("%Y-%m-%d %H:%M") if s.completion_date else "",
                    s.notes])

    # Payments sheet
    ws3 = wb.create_sheet("Payments")
    _style_header(ws3, ["ID", "Service ID", "Amount", "Method", "Date"])
    for p in Payment.query.filter_by(user_id=user_id).order_by(Payment.created_at).all():
        ws3.append([p.id, p.service_id, float(p.amount or 0), p.method,
                    p.created_at.strftime("%Y-%m-%d %H:%M") if p.created_at else ""])

    # Summary
    ws4 = wb.create_sheet("Summary")
    total_rev = sum(float(s.amount or 0) for s in services if s.status == "Completed")
    ws4.append(["Operator", current_user.full_name])
    ws4.append(["Email", current_user.email])
    ws4.append(["Shop", current_user.shop_name or ""])
    ws4.append(["Exported", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")])
    ws4.append([])
    ws4.append(["Total Customers", Customer.query.filter_by(user_id=user_id).count()])
    ws4.append(["Total Services", len(services)])
    ws4.append(["Completed Revenue (₹)", round(total_rev, 2)])
    for col in ("A", "B"):
        ws4.column_dimensions[col].width = 28
    for r in range(1, 5):
        ws4.cell(row=r, column=1).font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"cybercafe-history-{datetime.utcnow().strftime('%Y%m%d-%H%M')}.xlsx"
    return send_file(buf,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=fname)


# ---------------- Custom services CRUD (operator permission) ----------------

@export_bp.route("/my-services")
@login_required
def list_custom():
    items = (CustomService.query.filter_by(user_id=current_user.id)
             .order_by(CustomService.created_at.desc()).all())
    return render_template("service/custom_list.html", items=items)


def _own_custom(cs):
    if cs.user_id != current_user.id:
        abort(403)


@export_bp.route("/my-services/new", methods=["GET", "POST"])
@login_required
def new_custom():
    if request.method == "POST":
        name_en = (request.form.get("name_en") or "").strip()
        if not name_en:
            flash("Service name (English) is required.", "danger")
        else:
            cs = CustomService(
                user_id=current_user.id,
                name_en=name_en[:120],
                name_hi=(request.form.get("name_hi") or "").strip()[:120],
                category=(request.form.get("category") or "Custom").strip()[:60] or "Custom",
                emoji=(request.form.get("emoji") or "🛠️").strip()[:8] or "🛠️",
                fee=float(request.form.get("fee") or 0),
                is_active=True,
            )
            db.session.add(cs); db.session.commit()
            flash("Service added to your dashboard.", "success")
            return redirect(url_for("export.list_custom"))
    return render_template("service/custom_form.html", item=None, title="Add Service")


@export_bp.route("/my-services/<int:cid>/edit", methods=["GET", "POST"])
@login_required
def edit_custom(cid):
    cs = CustomService.query.get_or_404(cid); _own_custom(cs)
    if request.method == "POST":
        cs.name_en = (request.form.get("name_en") or cs.name_en).strip()[:120]
        cs.name_hi = (request.form.get("name_hi") or "").strip()[:120]
        cs.category = (request.form.get("category") or "Custom").strip()[:60] or "Custom"
        cs.emoji = (request.form.get("emoji") or "🛠️").strip()[:8] or "🛠️"
        try:
            cs.fee = float(request.form.get("fee") or 0)
        except ValueError:
            cs.fee = cs.fee
        cs.is_active = bool(request.form.get("is_active"))
        db.session.commit()
        flash("Service updated.", "success")
        return redirect(url_for("export.list_custom"))
    return render_template("service/custom_form.html", item=cs, title="Edit Service")


@export_bp.route("/my-services/<int:cid>/delete", methods=["POST"])
@login_required
def delete_custom(cid):
    cs = CustomService.query.get_or_404(cid); _own_custom(cs)
    db.session.delete(cs); db.session.commit()
    flash("Service removed.", "info")
    return redirect(url_for("export.list_custom"))
